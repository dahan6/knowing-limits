"""Load text, images and tables from common document formats."""
import base64
import hashlib
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Tuple


@dataclass
class DocumentPart:
    """A single piece of content extracted from a document."""

    type: str  # "text" | "image" | "table"
    source: str  # human-readable label, e.g. filename#page=1
    content: str  # text / markdown / ocr text / caption
    extra: Dict = field(default_factory=dict)


# ---------------------------------------------------------------------------
# Text helpers
# ---------------------------------------------------------------------------
def _read_text(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gbk", "gb2312", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError(f"无法解码文本文件: {path}")


# ---------------------------------------------------------------------------
# OCR / image caption (lazy singleton)
# ---------------------------------------------------------------------------
_ocr_engine = None


def _get_ocr():
    global _ocr_engine
    if _ocr_engine is None:
        try:
            from rapidocr import RapidOCR
        except ImportError as e:  # pragma: no cover
            raise ImportError("请安装 rapidocr: pip install rapidocr") from e
        _ocr_engine = RapidOCR()
    return _ocr_engine


def _ocr_image(image_path: Path) -> str:
    """Run OCR on an image and return recognized text."""
    engine = _get_ocr()
    output = engine(str(image_path))
    if output is None or not getattr(output, "txts", None):
        return ""
    texts = [t.strip() for t in output.txts if t and t.strip()]
    return "\n".join(texts)


# ---------------------------------------------------------------------------
# Markdown table helpers
# ---------------------------------------------------------------------------
def _rows_to_markdown(rows: List[List[str]]) -> str:
    if not rows:
        return ""
    # Normalize cells: strip whitespace, replace internal newlines
    cleaned = []
    for row in rows:
        cleaned.append([str(cell or "").replace("|", "\\|").replace("\n", " ").strip() for cell in row])
    lines = ["| " + " | ".join(cleaned[0]) + " |"]
    lines.append("| " + " | ".join(["---"] * len(cleaned[0])) + " |")
    for row in cleaned[1:]:
        # pad short rows
        row = row + [""] * (len(cleaned[0]) - len(row))
        lines.append("| " + " | ".join(row[: len(cleaned[0])]) + " |")
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def _load_pdf(
    path: Path,
    image_root: Optional[Path] = None,
) -> List[DocumentPart]:
    try:
        import fitz  # PyMuPDF
    except ImportError as e:  # pragma: no cover
        raise ImportError("请安装 pymupdf: pip install pymupdf") from e

    parts: List[DocumentPart] = []
    doc = fitz.open(str(path))
    base_name = path.name

    if image_root is not None:
        image_root.mkdir(parents=True, exist_ok=True)

    for page_idx in range(len(doc)):
        page = doc.load_page(page_idx)
        page_label = f"{base_name}#page={page_idx + 1}"

        # 1. text
        text = page.get_text()
        if text and text.strip():
            parts.append(DocumentPart(type="text", source=page_label, content=text.strip()))

        # 2. tables (via pdfplumber for better formatting)
        try:
            import pdfplumber
        except ImportError:  # pragma: no cover
            pdfplumber = None

        if pdfplumber is not None:
            with pdfplumber.open(str(path)) as plumber:
                if page_idx < len(plumber.pages):
                    tables = plumber.pages[page_idx].extract_tables()
                    for tidx, table in enumerate(tables or [], start=1):
                        if table:
                            md = _rows_to_markdown(table)
                            if md:
                                parts.append(
                                    DocumentPart(
                                        type="table",
                                        source=f"{base_name}#table=page={page_idx + 1},idx={tidx}",
                                        content=md,
                                    )
                                )

        # 3. images
        if image_root is not None:
            img_list = page.get_images(full=True)
            for img_idx, img in enumerate(img_list, start=1):
                xref = img[0]
                try:
                    pix = fitz.Pixmap(doc, xref)
                    if pix.n > 4:  # CMYK: convert to RGB
                        pix = fitz.Pixmap(fitz.csRGB, pix)
                    img_filename = f"{path.stem}_p{page_idx + 1}_img{img_idx}.png"
                    img_path = image_root / img_filename
                    pix.save(str(img_path))
                    pix = None
                    ocr_text = _ocr_image(img_path)
                    content = f"[图片内容]\n{ocr_text}".strip()
                    if content == "[图片内容]":
                        content = "[图片：未能识别出文字]"
                    parts.append(
                        DocumentPart(
                            type="image",
                            source=f"{base_name}#image=page={page_idx + 1},idx={img_idx}",
                            content=content,
                            extra={"image_path": str(img_path)},
                        )
                    )
                except Exception:
                    # Skip unprocessable images
                    continue

    return parts


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------
def _load_docx(path: Path) -> List[DocumentPart]:
    try:
        from docx import Document
    except ImportError as e:  # pragma: no cover
        raise ImportError("请安装 python-docx: pip install python-docx") from e

    parts: List[DocumentPart] = []
    doc = Document(str(path))

    # paragraphs
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    if paragraphs:
        parts.append(
            DocumentPart(
                type="text",
                source=path.name,
                content="\n".join(paragraphs),
            )
        )

    # tables
    for tidx, table in enumerate(doc.tables, start=1):
        rows = []
        for row in table.rows:
            rows.append([cell.text for cell in row.cells])
        md = _rows_to_markdown(rows)
        if md:
            parts.append(
                DocumentPart(
                    type="table",
                    source=f"{path.name}#table=idx={tidx}",
                    content=md,
                )
            )

    return parts


# ---------------------------------------------------------------------------
# XLSX / CSV
# ---------------------------------------------------------------------------
def _load_xlsx(path: Path) -> List[DocumentPart]:
    try:
        import openpyxl
    except ImportError as e:  # pragma: no cover
        raise ImportError("请安装 openpyxl: pip install openpyxl") from e

    parts: List[DocumentPart] = []
    wb = openpyxl.load_workbook(str(path), data_only=True)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append([str(cell) if cell is not None else "" for cell in row])
        md = _rows_to_markdown(rows)
        if md:
            parts.append(
                DocumentPart(
                    type="table",
                    source=f"{path.name}#sheet={sheet_name}",
                    content=md,
                )
            )
    return parts


def _load_csv(path: Path) -> List[DocumentPart]:
    import csv

    rows = []
    with open(path, "r", encoding="utf-8", errors="ignore", newline="") as f:
        reader = csv.reader(f)
        for row in reader:
            rows.append(row)
    md = _rows_to_markdown(rows)
    if md:
        return [
            DocumentPart(
                type="table",
                source=path.name,
                content=md,
            )
        ]
    return []


# ---------------------------------------------------------------------------
# Image files
# ---------------------------------------------------------------------------
def _load_image(path: Path) -> List[DocumentPart]:
    ocr_text = _ocr_image(path)
    content = f"[图片内容]\n{ocr_text}".strip()
    if content == "[图片内容]":
        content = "[图片：未能识别出文字]"
    return [
        DocumentPart(
            type="image",
            source=path.name,
            content=content,
            extra={"image_path": str(path)},
        )
    ]


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------
SUPPORTED_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".pdf",
    ".docx",
    ".csv",
    ".xlsx",
    ".json",
    ".py",
    ".js",
    ".ts",
    ".html",
    ".xml",
}
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif"}


def load_document(
    path: Path,
    image_root: Optional[Path] = None,
) -> List[DocumentPart]:
    """
    Load a document and return a list of DocumentPart objects.

    For PDFs, extracted images are saved under ``image_root`` (if provided)
    and their OCR text is stored in the part content.
    """
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _load_pdf(path, image_root=image_root)
    elif suffix == ".docx":
        return _load_docx(path)
    elif suffix == ".xlsx":
        return _load_xlsx(path)
    elif suffix == ".csv":
        return _load_csv(path)
    elif suffix in IMAGE_EXTENSIONS:
        return _load_image(path)
    elif suffix in SUPPORTED_EXTENSIONS:
        return [DocumentPart(type="text", source=path.name, content=_read_text(path))]
    else:
        # Fallback: try as plain text
        return [DocumentPart(type="text", source=path.name, content=_read_text(path))]
